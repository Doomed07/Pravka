import json
import os

import pytest
from openpyxl import load_workbook

from pravki import journal
from pravki.cli import main
from pravki.llm import Explainer, LLMError, OpenAICompatClient, check_output, is_local
from pravki.model import Change
from pravki.pipeline import DISCLAIMER, review_files
from pravki.report import render_html, render_xlsx
from pravki.synth import generate_case
from tests.helpers import dele, ins, make_docx, p, r

PB = os.path.join(os.path.dirname(os.path.dirname(__file__)), "playbooks", "services_example.yaml")


@pytest.fixture()
def case(tmp_path):
    return generate_case(str(tmp_path / "case"), "services", 7)


def test_cli_review_writes_reports(case, tmp_path, capsys):
    out = tmp_path / "out"
    rc = main(["review", case["paths"]["tracked"], "--base", case["paths"]["base"], "--playbook", PB, "--out", str(out)])
    assert rc == 0
    for f in ("report.html", "report.xlsx", "changes.json"):
        assert (out / f).exists()
    assert DISCLAIMER in capsys.readouterr().out
    data = json.loads((out / "changes.json").read_text(encoding="utf-8"))
    assert data["disclaimer"] == DISCLAIMER and data["changes"]
    html = (out / "report.html").read_text(encoding="utf-8")
    assert DISCLAIMER in html and "<ins>" in html and "<del>" in html


def test_report_ends_with_disclaimer_everywhere(case, tmp_path):
    rv = review_files(case["paths"]["tracked"], case["paths"]["base"], PB)
    render_xlsx(rv, str(tmp_path / "r.xlsx"))
    wb = load_workbook(tmp_path / "r.xlsx")
    assert set(wb.sheetnames) == {"Правки", "Сводка", "Для бизнеса"}
    for ws in wb.worksheets:
        texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert any(DISCLAIMER in t for t in texts), ws.title


def test_html_escapes_untrusted_contract_text(tmp_path):
    evil = "<script>alert(1)</script>"
    t = make_docx(tmp_path / "t.docx", p(r("1.1. Срок "), dele("10"), ins("30" + evil.replace("<", "&lt;").replace(">", "&gt;")), r(" дней.")))
    rv = review_files(str(t))
    html = render_html(rv)
    assert "<script>" not in html and "&lt;script&gt;" in html


def test_xlsx_neutralises_formulas(tmp_path):
    t = make_docx(tmp_path / "t.docx", p(r("1.1. Сумма "), ins("=HYPERLINK(&quot;http://evil&quot;,&quot;x&quot;)"), r(".")))
    rv = review_files(str(t))
    render_xlsx(rv, str(tmp_path / "r.xlsx"))
    ws = load_workbook(tmp_path / "r.xlsx")["Правки"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                assert not c.value.startswith("=")


def test_priority_topics_come_first(tmp_path):
    body = p(r("1.1. Срок оплаты "), dele("10"), ins("30"), r(" дней.")) + \
        p(r("1.2. Каждая из сторон вправе отказаться от договора, уведомив за "), dele("30"), ins("5"), r(" календарных дней."))
    t = make_docx(tmp_path / "t.docx", body)
    rv = review_files(str(t), None, PB)
    assert rv.changes[0].topic == "расторжение и отказ от договора"


def test_journal_marks_repeated_rejected_edit(case, tmp_path):
    proj = str(tmp_path / "proj")
    rv1 = review_files(case["paths"]["tracked"], case["paths"]["base"], PB)
    journal.init_project(proj, "Тест")
    journal.add_round(proj, rv1)
    first = next(c for c in rv1.changes if c.materiality != "техническая")
    journal.decide(proj, first.id, "rejected", "не согласны")
    rv2 = review_files(case["paths"]["tracked"], case["paths"]["base"], PB)
    n = journal.add_round(proj, rv2)
    assert n == 2
    again = next(c for c in rv2.changes if c.fingerprint() == first.fingerprint())
    assert again.status == "repeated" and "отклоня" in again.repeat_note
    assert "Раунд 2" in journal.history_text(proj)


def test_journal_errors(tmp_path):
    with pytest.raises(journal.JournalError):
        journal.load(str(tmp_path / "нет"))
    journal.init_project(str(tmp_path / "p"), "x")
    with pytest.raises(journal.JournalError):
        journal.init_project(str(tmp_path / "p"), "x")
    with pytest.raises(journal.JournalError):
        journal.decide(str(tmp_path / "p"), "E001", "accepted")
    with pytest.raises(journal.JournalError):
        journal.decide(str(tmp_path / "p"), "E001", "maybe")


def test_cli_errors_are_readable(tmp_path, capsys):
    assert main(["review", str(tmp_path / "нет.docx")]) == 2
    assert "не найден" in capsys.readouterr().err
    bad = tmp_path / "bad.docx"
    bad.write_text("не docx")
    assert main(["review", str(bad)]) == 2
    assert "Ошибка" in capsys.readouterr().err


def test_cli_project_flow(case, tmp_path, capsys):
    proj = str(tmp_path / "proj")
    assert main(["review", case["paths"]["tracked"], "--base", case["paths"]["base"], "--playbook", PB,
                 "--project", proj, "--name", "Демо"]) == 0
    assert os.path.exists(os.path.join(proj, "rounds", "01", "report.html"))
    assert main(["decide", proj, "E001", "counter", "--note", "встречная"]) == 0
    assert main(["history", proj]) == 0
    assert "встречная" in capsys.readouterr().out


# ---- слой модели ----
class FakeClient:
    def __init__(self, answer):
        self.answer = answer

    def complete(self, system, user):
        return self.answer


def _change():
    ch = Change(before="0,1%", after="0,5%", context="Неустойка 0,1% в день", context_after="Неустойка 0,5% в день",
                explanation="Ставка выше допустимой (не более 0,3%).", materiality="высокая", rule_id="SV-01")
    ch.detail["norms"] = ["ст. 330 ГК РФ"]
    return ch


def test_llm_answer_with_invented_number_is_rejected():
    ch = _change()
    Explainer(FakeClient("Ставка выросла до 0,5%, а по закону допустимо 7%.")).apply(ch)
    assert ch.explanation.startswith("Ставка выше") and any("отброшен" in n for n in ch.notes)


def test_llm_answer_with_unknown_article_is_rejected():
    ch = _change()
    Explainer(FakeClient("Ставка 0,5% противоречит ст. 999 ГК РФ.")).apply(ch)
    assert ch.explanation.startswith("Ставка выше")


def test_llm_good_answer_is_used():
    ch = _change()
    e = Explainer(FakeClient("Контрагент поднял ставку неустойки с 0,1% до 0,5%: это выше допустимого."))
    e.apply(ch)
    assert ch.explanation.startswith("Контрагент поднял") and e.used == 1
    assert check_output("Ставка 0,5% по ст. 330 ГК РФ", ch, ["ст. 330 ГК РФ"]) is None


def test_llm_local_only_by_default():
    assert is_local("http://localhost:8080/v1") and is_local("http://192.168.1.5:11434/v1")
    assert not is_local("https://api.example.com/v1")
    with pytest.raises(LLMError):
        OpenAICompatClient("https://api.example.com/v1", "m")
    OpenAICompatClient("https://api.example.com/v1", "m", allow_remote=True)


def test_llm_unavailable_keeps_template(case):
    class Broken:
        def complete(self, s, u):
            raise LLMError("нет соединения")

    rv = review_files(case["paths"]["tracked"], case["paths"]["base"], PB, explainer=Explainer(Broken()))
    assert rv.changes and all(c.explanation or c.materiality == "техническая" for c in rv.changes)
