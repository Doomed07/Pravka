"""Отчёты: HTML (для чтения) и Excel (для работы).

Оформление правок — как принято в нашей редакции: вставки подчёркнуты, удаления зачёркнуты,
комментарии вынесены на поля. Весь текст договора недоверенный, поэтому в HTML он экранируется.
"""
from __future__ import annotations

import html
import json
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment as XlComment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .model import Change
from .pipeline import DISCLAIMER, Review, file_label
from .playbook import PRIORITY_DEFAULT

SOURCE_RU = {
    "tracked": "в рецензировании",
    "silent": "ТИХАЯ ПРАВКА (без отметки)",
    "diff": "найдена сравнением",
}
KIND_RU = {
    "insert": "вставка", "delete": "удаление", "replace": "замена", "move": "перемещение",
    "para_insert": "новый пункт", "para_delete": "удалён пункт", "comment": "комментарий",
}
LEVEL_RU = {"accept": "принять", "fallback": "уступка", "reject": "отклонить / встречная", "review": "решает юрист"}
LEVEL_CSS = {"accept": "ok", "fallback": "fb", "reject": "bad", "review": "rev"}
ORDER = {"reject": 0, "review": 1, "fallback": 2, "accept": 3}


def _e(s: str) -> str:
    return html.escape(s or "", quote=True)


def redline_html(ch: Change) -> str:
    """Абзац с правкой в тексте: удалённое зачёркнуто, добавленное подчёркнуто."""
    if ch.kind == "comment":
        return f'<span class="anchor">{_e(ch.before) or "—"}</span>'
    if ch.kind == "para_insert":
        return f'<ins>{_e(ch.after)}</ins>'
    if ch.kind == "para_delete":
        return f'<del>{_e(ch.before)}</del>'
    if ch.kind == "move":
        return f'<del>{_e(ch.before)}</del> <span class="muted">(перемещено: п. {_e(str(ch.detail.get("move_from", "")))} → п. {_e(str(ch.detail.get("move_to", "")))})</span>'
    base = ch.context_after if ch.after else ch.context
    frag = ch.after if ch.after else ch.before
    i = base.find(frag) if (base and frag) else -1
    if i < 0:
        out = ""
        if ch.before:
            out += f"<del>{_e(ch.before)}</del> "
        if ch.after:
            out += f"<ins>{_e(ch.after)}</ins>"
        return out
    pre, post = base[:i], base[i + len(frag):]
    mid = ""
    if ch.before and ch.after:
        mid = f"<del>{_e(ch.before)}</del> <ins>{_e(ch.after)}</ins>"
    elif ch.after:
        mid = f"<ins>{_e(ch.after)}</ins>"
    else:
        mid = f"<del>{_e(ch.before)}</del>"
    return _e(pre) + mid + _e(post)


CSS = """
:root{--bg:#fff;--fg:#1d2127;--muted:#667085;--line:#e4e7ec;--card:#f9fafb;--ins:#0b57d0;--del:#c5221f;
--ok:#e6f4ea;--fb:#fef3e0;--bad:#fce8e6;--rev:#fff8c5;--accent:#2f4858}
@media (prefers-color-scheme:dark){:root{--bg:#14171c;--fg:#e8eaed;--muted:#9aa4b2;--line:#2b313a;--card:#1b2027;
--ins:#8ab4f8;--del:#f28b82;--ok:#173a22;--fb:#3f3010;--bad:#4a1d1b;--rev:#3c3a12;--accent:#9db4c0}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.5 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
main{max-width:1180px;margin:0 auto;padding:24px 16px 64px}
h1{font-size:22px;margin:0 0 4px}h2{font-size:17px;margin:32px 0 12px;border-bottom:1px solid var(--line);padding-bottom:6px}
.muted{color:var(--muted)}.meta{font-size:13px;color:var(--muted);margin-bottom:16px}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:8px;margin:12px 0}
.stat{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:8px 12px}.stat b{display:block;font-size:20px}
.warn{background:var(--rev);border-left:4px solid #e0a800;padding:8px 12px;margin:8px 0;border-radius:4px}
.item{display:grid;grid-template-columns:minmax(0,3fr) minmax(220px,1.4fr);gap:0 16px;border:1px solid var(--line);
border-radius:10px;margin:12px 0;background:var(--card);overflow:hidden}
.item.silent{border-color:#e0a800;box-shadow:0 0 0 1px #e0a800}
.main{padding:12px 14px;min-width:0}.margin{padding:12px 14px;border-left:1px dashed var(--line);font-size:13.5px;min-width:0}
.head{display:flex;flex-wrap:wrap;gap:6px 10px;align-items:baseline;margin-bottom:6px;font-size:13px}
.num{font-weight:700}.tag{border:1px solid var(--line);border-radius:999px;padding:0 8px;font-size:12px;color:var(--muted)}
.tag.silent{background:var(--rev);color:var(--fg);border-color:#e0a800}
.text{white-space:pre-wrap;overflow-wrap:anywhere}ins{color:var(--ins);text-decoration:underline;text-decoration-thickness:2px}
del{color:var(--del);text-decoration:line-through}.anchor{background:var(--rev)}
.badge{display:inline-block;border-radius:6px;padding:1px 8px;font-weight:600;font-size:13px}
.ok{background:var(--ok)}.fb{background:var(--fb)}.bad{background:var(--bad)}.rev{background:var(--rev)}
.margin p{margin:0 0 8px}.margin .lbl{display:block;font-size:11.5px;letter-spacing:.03em;text-transform:uppercase;color:var(--muted)}
.cmt{border-left:3px solid var(--accent);padding-left:8px;margin-bottom:8px}
details summary{cursor:pointer;color:var(--muted)}
footer{margin-top:40px;padding-top:12px;border-top:1px solid var(--line);color:var(--muted);font-size:13px}
@media (max-width:760px){.item{grid-template-columns:1fr}.margin{border-left:0;border-top:1px dashed var(--line)}}
@media print{.item{break-inside:avoid}}
"""


def _item_html(ch: Change) -> str:
    src = SOURCE_RU.get(ch.source, ch.source)
    cls = "item silent" if ch.source == "silent" else "item"
    tags = [f'<span class="tag{" silent" if ch.source == "silent" else ""}">{_e(src)}</span>',
            f'<span class="tag">{_e(KIND_RU.get(ch.kind, ch.kind))}</span>']
    if ch.author:
        tags.append(f'<span class="tag">{_e(ch.author)}</span>')
    if ch.status != "new":
        tags.append(f'<span class="tag silent">{_e(ch.status)}</span>')
    clause = f"п. {_e(ch.clause)}" if ch.clause else "пункт не определён"
    head = f'<div class="head"><span class="num">{_e(ch.id)} · {clause}</span>' + "".join(tags)
    if ch.heading:
        head += f'<span class="muted">{_e(ch.heading)}</span>'
    head += "</div>"
    main = f'{head}<div class="text">{redline_html(ch)}</div>'

    m = []
    if ch.kind == "comment":
        m.append(f'<div class="cmt"><span class="lbl">Комментарий контрагента</span>{_e(ch.after)}</div>')
    if ch.level:
        m.append(f'<p><span class="badge {LEVEL_CSS.get(ch.level, "rev")}">{_e(ch.action)}</span></p>')
    meta = []
    if ch.topic:
        meta.append(f"<b>Тема:</b> {_e(ch.topic)}")
    if ch.areas:
        meta.append(f"<b>Области:</b> {_e(ch.area)}")
    if ch.materiality:
        meta.append(f"<b>Существенность:</b> {_e(ch.materiality)}")
    if ch.rule_id:
        meta.append(f"<b>Правило:</b> {_e(ch.rule_id)} {_e(ch.rule_title)}")
    if ch.flags:
        meta.append("<b>Признаки:</b> " + _e(", ".join(ch.flags)))
    if meta:
        m.append("<p>" + "<br>".join(meta) + "</p>")
    if ch.explanation:
        m.append(f'<p><span class="lbl">Пояснение</span>{_e(ch.explanation)}</p>')
    if ch.suggested and ch.level in ("reject", "fallback", "review"):
        m.append(f'<p><span class="lbl">Позиция по плейбуку (черновик)</span>{_e(ch.suggested)}</p>')
    if ch.escalate_to:
        m.append(f'<p><span class="lbl">Согласовать с</span>{_e(ch.escalate_to)}</p>')
    for n in ch.notes:
        m.append(f'<p class="muted">{_e(n)}</p>')
    if ch.repeat_note:
        m.append(f'<p class="muted">{_e(ch.repeat_note)}</p>')
    return f'<article class="{cls}"><div class="main">{main}</div><div class="margin">{"".join(m)}</div></article>'


def business_summary(rv: Review) -> list[str]:
    """Короткая сводка для бизнеса: что меняется по деньгам, срокам и обязательствам, без юридической стратегии."""
    lines = []
    for ch in sorted(rv.changes, key=lambda c: c.order):
        if ch.materiality == "техническая" or ch.kind == "comment":
            continue
        if not set(ch.areas) & {"коммерческая", "финансовая", "бухгалтерская и налоговая"}:
            continue
        what = ch.topic or "условие"
        if ch.kind in ("para_insert",):
            txt = f"добавлен пункт: «{_short(ch.after)}»"
        elif ch.kind in ("para_delete",):
            txt = f"удалён пункт: «{_short(ch.before)}»"
        else:
            txt = f"было «{_short(ch.before) or '—'}» → стало «{_short(ch.after) or '—'}»"
        mark = " [без отметки в рецензировании]" if ch.source == "silent" else ""
        lines.append(f"п. {ch.clause or '?'} · {what}: {txt}{mark}")
    return lines


def _short(s: str, n: int = 110) -> str:
    s = " ".join((s or "").split())
    return s if len(s) <= n else s[: n - 1] + "…"


def render_html(rv: Review, priority_topics: list[str] | None = None) -> str:
    prio = priority_topics or PRIORITY_DEFAULT
    pr = [c for c in rv.changes if c.topic in prio and c.materiality != "техническая"]
    tech = [c for c in rv.changes if c.materiality == "техническая" and c.kind != "comment"]
    rest = [c for c in rv.changes if c not in pr and c not in tech]
    counts = rv.counts()
    area_counter: Counter = Counter()
    for c in rv.changes:
        if c.materiality in ("высокая", "средняя"):
            for a in c.areas:
                if a != "юридическая":
                    area_counter[a] += 1

    out = [f"<!doctype html><html lang='ru'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>",
           f"<title>Разбор правок — {_e(file_label(rv.theirs_path))}</title><style>{CSS}</style></head><body><main>",
           f"<h1>Разбор правок контрагента</h1>",
           f"<div class='meta'>Версия контрагента: {_e(file_label(rv.theirs_path))} · "
           f"Наша отправленная версия: {_e(file_label(rv.base_path)) or 'не указана'} · "
           f"Плейбук: {_e(rv.playbook_name) or 'не указан'}{' · мы — ' + _e(rv.our_role) if rv.our_role else ''} · {_e(rv.created)}</div>"]
    for w in rv.warnings:
        out.append(f"<div class='warn'>{_e(w)}</div>")
    out.append("<div class='stats'>" + "".join(
        f"<div class='stat'><b>{v}</b>{_e(k)}</div>" for k, v in counts.items()) + "</div>")
    if area_counter:
        out.append("<p class='muted'>Кого подключить (существенные правки вне юридической области): " +
                   _e(", ".join(f"{a} — {n}" for a, n in area_counter.most_common())) + "</p>")

    if pr:
        out.append("<h2>Сначала: возмещение убытков, ограничение ответственности, расторжение</h2>")
        out.extend(_item_html(c) for c in pr)
    if rest:
        out.append("<h2>Остальные правки</h2>")
        out.extend(_item_html(c) for c in rest)
    if tech:
        out.append(f"<h2>Технические правки ({len(tech)})</h2><details><summary>Опечатки, знаки, регистр, пробелы — показать</summary>")
        out.extend(_item_html(c) for c in tech)
        out.append("</details>")
    if not rv.changes:
        out.append("<p>Правок не найдено.</p>")

    bs = business_summary(rv)
    if bs:
        out.append("<h2>Для бизнеса: что меняется по деньгам, срокам и обязательствам</h2><ul>" +
                   "".join(f"<li>{_e(x)}</li>" for x in bs) + "</ul>")
    out.append(f"<footer>{_e(DISCLAIMER)}<br>Автоматический разбор может ошибаться: решение по каждой правке принимает юрист.</footer>")
    out.append("</main></body></html>")
    return "".join(out)


FILL = {
    "reject": PatternFill("solid", fgColor="F8D7DA"),
    "review": PatternFill("solid", fgColor="FFF3B0"),
    "fallback": PatternFill("solid", fgColor="FFE5B4"),
    "accept": PatternFill("solid", fgColor="D4EDDA"),
}
HEAD_FILL = PatternFill("solid", fgColor="2F4858")


def render_xlsx(rv: Review, path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Правки"
    cols = [
        ("№", 7), ("Пункт", 9), ("Раздел", 22), ("Источник", 16), ("Автор", 16), ("Тип", 12),
        ("Было (зачёркнуто)", 38), ("Стало (подчёркнуто)", 38), ("Комментарий контрагента", 30),
        ("Тема", 24), ("Области компетенции", 24), ("Существенность", 14), ("Правило плейбука", 26),
        ("Рекомендация", 24), ("Пояснение", 50), ("Позиция по плейбуку (черновик)", 36), ("Согласовать с", 22),
        ("Нужно решение юриста", 12), ("Вне плейбука", 10), ("Статус", 12), ("Признаки", 20),
    ]
    for i, (name, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"

    prio = PRIORITY_DEFAULT
    rows = sorted(rv.changes, key=lambda c: (0 if (c.topic in prio and c.materiality != "техническая") else 1,
                                             ORDER.get(c.level, 1), c.order))
    for r, ch in enumerate(rows, 2):
        comment_text = ch.after if ch.kind == "comment" else ""
        vals = [
            ch.id, ch.clause, ch.heading, SOURCE_RU.get(ch.source, ch.source), ch.author, KIND_RU.get(ch.kind, ch.kind),
            "" if ch.kind == "comment" else ch.before, "" if ch.kind == "comment" else ch.after, comment_text,
            ch.topic, ch.area, ch.materiality, f"{ch.rule_id} {ch.rule_title}".strip(), ch.action, ch.explanation,
            ch.suggested if ch.level in ("reject", "fallback", "review") else "", ch.escalate_to,
            "да" if ch.lawyer_decision_needed else "", "да" if ch.out_of_playbook else "", ch.status, ", ".join(ch.flags),
        ]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=_cell(v))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=7).font = Font(strike=True, color="C5221F")
        ws.cell(row=r, column=8).font = Font(underline="single", color="0B57D0")
        if ch.level in FILL:
            ws.cell(row=r, column=14).fill = FILL[ch.level]
        if ch.source == "silent":
            ws.cell(row=r, column=4).fill = FILL["review"]
            ws.cell(row=r, column=4).font = Font(bold=True)
        if ch.notes or ch.repeat_note:
            ws.cell(row=r, column=1).comment = XlComment("\n".join(ch.notes + ([ch.repeat_note] if ch.repeat_note else [])), "Разбор")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(rows) + 1)}"
    end = len(rows) + 3
    ws.cell(row=end, column=1, value=DISCLAIMER).font = Font(italic=True, color="667085")

    s = wb.create_sheet("Сводка")
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 14
    s.append(["Файл контрагента", file_label(rv.theirs_path)])
    s.append(["Отправленная нами версия", file_label(rv.base_path) or "не указана"])
    s.append(["Плейбук", rv.playbook_name or "не указан"])
    s.append(["Мы", rv.our_role or ""])
    s.append(["Дата разбора", rv.created])
    s.append([])
    for k, v in rv.counts().items():
        s.append([k, v])
    s.append([])
    s.append(["По областям (существенные правки)", ""])
    by_area: dict[str, int] = defaultdict(int)
    for c in rv.changes:
        if c.materiality in ("высокая", "средняя"):
            for a in c.areas:
                by_area[a] += 1
    for a, n in sorted(by_area.items(), key=lambda kv: -kv[1]):
        s.append([a, n])
    s.append([])
    s.append(["По темам", ""])
    by_topic = Counter(c.topic or "—" for c in rv.changes)
    for t, n in by_topic.most_common():
        s.append([t, n])
    for w in rv.warnings:
        s.append([])
        s.append([w])
    s.append([])
    s.append([DISCLAIMER])
    for row in s.iter_rows(min_row=1, max_row=5):
        row[0].font = Font(bold=True)

    b = wb.create_sheet("Для бизнеса")
    b.column_dimensions["A"].width = 140
    b.append(["Что меняется по деньгам, срокам и обязательствам"])
    b["A1"].font = Font(bold=True)
    for line in business_summary(rv) or ["Существенных коммерческих изменений не найдено."]:
        b.append([line])
        b.cell(row=b.max_row, column=1).alignment = Alignment(wrap_text=True)
    b.append([])
    b.append([DISCLAIMER])
    wb.save(path)


def _cell(v):
    """Значения из чужого документа не должны превращаться в формулы Excel."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def write_json(rv: Review, path: str) -> None:
    data = {
        "disclaimer": DISCLAIMER,
        "theirs": file_label(rv.theirs_path),
        "base": file_label(rv.base_path),
        "playbook": rv.playbook_name,
        "our_role": rv.our_role,
        "mode": rv.mode,
        "created": rv.created,
        "warnings": rv.warnings,
        "counts": rv.counts(),
        "changes": [c.to_dict() for c in rv.changes],
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
