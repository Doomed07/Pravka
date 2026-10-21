"""Анкета для разметки: правки контрагента → xlsx, где юрист отмечает тему и решение.

Зачем. Качество программы (тема, решение) проверяется только на правках, по которым юрист сказал, как правильно.
Анкета собирает правки из разбора и оставляет рядом с каждой свободные поля для ответа; ответы читаются обратно
в «золотой набор» (`load_labels`). Программа лишь предзаполняет то, что поняла сама, чтобы юристу было что исправлять.

Шум не раздуваем: десять «склеек абзацев» в реквизитах — одна строка, а не десять; две вставки в одном абзаце,
которые вместе дают одну мысль («, но не позднее» + «апреля 2026»), — тоже одна. Идентификаторы правок строки сохраняются
(колонка «ID»), и оценка по ответу относится ко всем им.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .classify import OTHER, TOPICS
from .model import Change
from .pipeline import DISCLAIMER


KIND_RU = {
    "insert": "вставка", "delete": "удаление", "replace": "замена", "move": "перенос",
    "para_insert": "новый абзац", "para_delete": "удалён абзац", "comment": "комментарий",
}
DECISIONS = [
    "Принять",
    "Уступить (в пределах позиции)",
    "Отклонить",
    "Встречное предложение",
    "Эскалировать (решает руководитель)",
    "Не требует решения (согласие, шум)",
]
TOPIC_OK = ["Да", "Нет"]
MAX_CELL = 700
SHORT = 60  # символов до+после: короче — считаем обрывком

SHEET_ROWS = "Разметка"
SHEET_GUIDE = "Памятка"
SHEET_CONTRACTS = "Договоры"
HEADERS = [
    "№", "Договор", "Пункт", "Заголовок пункта", "Что сделал контрагент", "Было", "Стало", "Абзац целиком (для контекста)",
    "Признаки программы", "Тема по программе", "Тема верна?", "Правильная тема", "Ваше решение", "Почему / граница уступки",
    "ID правок",
]
WIDTHS = [5, 18, 8, 26, 14, 42, 42, 48, 22, 22, 11, 24, 26, 44, 14]
INPUT_COLS = (11, 12, 13, 14)  # K..N: сюда пишет юрист


@dataclass
class Row:
    contract: str
    clause: str
    heading: str
    kinds: list[str]
    before: str
    after: str
    context: str
    flags: list[str]
    topic: str
    ids: list[str] = field(default_factory=list)
    grouped: int = 1


def _clip(s: str, n: int = MAX_CELL) -> str:
    s = (s or "").replace("\x00", "").strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


def _is_noise(c: Change) -> bool:
    return "склейка абзацев" in c.flags


def _short(c: Change) -> bool:
    return len(c.before or "") + len(c.after or "") <= SHORT


def _absorb(r: Row, c: Change) -> None:
    r.ids.append(c.id)
    r.grouped += 1
    if c.before:
        r.before = f"{r.before} | {c.before}".strip(" |")
    if c.after:
        r.after = f"{r.after} | {c.after}".strip(" |")
    kind = KIND_RU.get(c.kind, c.kind)
    if kind not in r.kinds:
        r.kinds.append(kind)
    for fl in c.flags + ([] if c.author else ["найдена сравнением"]):
        if fl not in r.flags:
            r.flags.append(fl)


def group_changes(changes: list[Change], contract: str, theirs_only: bool = True) -> list[Row]:
    """Свести правки в строки анкеты. Порядок — как в документе.

    Склеиваем только то, что юрист всё равно решает одним ответом:
    * «склейки абзацев» одного пункта — одна техническая строка;
    * короткие обрывки одного абзаца и «тихие» расхождения, найденные сравнением (", но не позднее" + "апреля 2026");
    * подряд идущие мелкие правки одного пункта вне тем (реквизиты: телефон, офис).
    Длинные правки одного абзаца остаются раздельными: это разные мысли («ущерб только реальный» и «ответственность
    прекращается по истечении срока приёмки»), решение по ним может быть разным.
    """
    items = [c for c in changes if not theirs_only or c.author != "Мы"]
    items.sort(key=lambda c: (c.order, c.id))
    rows: list[Row] = []
    noise: dict[str, Row] = {}
    by_par: dict[int, Row] = {}
    last_small: Row | None = None
    for c in items:
        if _is_noise(c):
            key = c.clause or "-"
            r = noise.get(key)
            if r is None:
                r = Row(contract, c.clause, c.heading, ["склейка абзацев"], "", "", c.context, ["склейка абзацев"], c.topic, grouped=0)
                noise[key] = r
                rows.append(r)
            r.ids.append(c.id)
            r.grouped += 1
            continue
        small = c.kind != "comment" and c.topic == OTHER and _short(c)
        silent = not c.author
        if c.kind != "comment" and c.paragraph_index >= 0:
            r = by_par.get(c.paragraph_index)
            if r is not None and (_short(c) or silent):
                _absorb(r, c)
                continue
        if small and last_small is not None and last_small.clause == c.clause:
            _absorb(last_small, c)
            if c.paragraph_index >= 0:
                by_par.setdefault(c.paragraph_index, last_small)
            continue
        if c.kind == "comment":
            before, after = f"(к тексту) {_clip(c.before, 260)}", c.after
        else:
            before, after = c.before, c.after
        r = Row(contract, c.clause, c.heading, [KIND_RU.get(c.kind, c.kind)], before, after, c.context or c.before,
                list(c.flags), c.topic, [c.id])
        if silent:
            r.flags.append("найдена сравнением")
        rows.append(r)
        if c.kind != "comment" and c.paragraph_index >= 0:
            by_par.setdefault(c.paragraph_index, r)
        last_small = r if small else None
    return rows


def _put(ws, row: int, col: int, value):
    """Записать значение так, чтобы текст из чужого договора не стал формулой Excel («=HYPERLINK(...)»)."""
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        c.data_type = "s"
    return c


def _style(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_worksheet(rows: list[Row], path: str, contracts: list[dict] | None = None, topics: list[str] | None = None) -> int:
    """Записать анкету. contracts: [{'name', 'kind', 'parties', 'note'}] — на листе «Договоры» (колонка «Ваша роль» пустая)."""
    topics = topics or [t.name for t in TOPICS] + [OTHER]
    wb = Workbook()
    f_base = Font(name="Arial", size=10)
    f_bold = Font(name="Arial", size=10, bold=True)
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_head = PatternFill("solid", fgColor="305496")
    fill_in = PatternFill("solid", fgColor="FFF2CC")
    fill_noise = PatternFill("solid", fgColor="EDEDED")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    # --- Памятка
    g = wb.active
    g.title = SHEET_GUIDE
    lines = [
        ("Анкета разметки правок контрагента", f_bold),
        ("Тексты обезличены (названия, ФИО, реквизиты заменены); смысл условий сохранён. Подтверждать в договоре ничего не нужно.", f_base),
        ("", f_base),
        ("Что заполнить (жёлтые ячейки на листе «Разметка», колонки K–N):", f_bold),
        ("K «Тема верна?» — Да, если тема по программе (колонка J) совпала с вашей.", f_base),
        ("L «Правильная тема» — только если в K «Нет»: выберите из списка.", f_base),
        ("M «Ваше решение» — как бы вы поступили с этой правкой. Списком: принять / уступить / отклонить / встречное предложение / эскалировать / не требует решения.", f_base),
        ("N «Почему / граница уступки» — одной строкой: почему так и до какого значения готовы уступить (срок, процент, сумма).", f_base),
        ("Лист «Договоры»: в колонке «Ваша роль» укажите, за какую сторону вы в этом договоре (Заказчик, Исполнитель, Покупатель…).", f_base),
        ("", f_base),
        ("Если правка непонятна без соседнего текста — см. колонку H («Абзац целиком»). Строка «склейка абзацев» — техническая правка вёрстки, решение не нужно.", f_base),
        ("Одна строка может объединять несколько правок одного абзаца (колонка O — их ID): ответ относится ко всем.", f_base),
        ("", f_base),
        ("Пример заполненной строки (его нет на листе «Разметка»):", f_bold),
        ("Пункт 5.2 · вставка · «Штраф не более 5% от суммы» · тема по программе «неустойка и штрафы» → Тема верна? Да · "
         "Решение: Уступить · Почему: 5% приемлемо, ниже 3% — эскалировать.", f_base),
        ("", f_base),
        (DISCLAIMER, f_bold),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = g.cell(row=i, column=1, value=text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width = 130

    # --- Договоры
    cs = wb.create_sheet(SHEET_CONTRACTS)
    for j, h in enumerate(["Договор (условное имя)", "Вид", "Стороны по тексту (по программе)", "Что за раунд", "Ваша роль"], 1):
        c = cs.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border, c.alignment = f_head, fill_head, border, wrap
    for i, k in enumerate(contracts or [], 2):
        vals = [k.get("name", ""), k.get("kind", ""), k.get("parties", ""), k.get("note", ""), ""]
        for j, v in enumerate(vals, 1):
            c = _put(cs, i, j, v)
            c.font, c.border, c.alignment = f_base, border, wrap
        cs.cell(row=i, column=5).fill = fill_in
    _style(cs, [34, 22, 60, 50, 24])
    cs.cell(row=(len(contracts or []) + 3), column=1, value=DISCLAIMER).font = f_bold

    # --- Разметка
    ws = wb.create_sheet(SHEET_ROWS)
    for j, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = f_head, fill_head, border
        c.alignment = Alignment(wrap_text=True, vertical="center")
    _style(ws, WIDTHS)
    ws.row_dimensions[1].height = 32
    for i, r in enumerate(rows, 2):
        flags = ", ".join(r.flags)
        if r.grouped > 1 and "склейка абзацев" in r.flags:
            flags = f"склейка абзацев × {r.grouped}"
        vals = [i - 1, r.contract, r.clause, _clip(r.heading, 120), ", ".join(r.kinds), _clip(r.before), _clip(r.after),
                _clip(r.context, 500), flags, r.topic, "", "", "", "", ", ".join(r.ids)]
        for j, v in enumerate(vals, 1):
            c = _put(ws, i, j, v)
            c.font, c.border, c.alignment = f_base, border, wrap
            if j in INPUT_COLS:
                c.fill = fill_in
            elif "склейка абзацев" in r.flags and j not in INPUT_COLS:
                c.fill = fill_noise
        longest = max(len(str(vals[5])) / 42, len(str(vals[6])) / 42, len(str(vals[7])) / 48, 1)
        ws.row_dimensions[i].height = min(409, 14 * (int(longest) + 2))
    last = len(rows) + 1
    ws.freeze_panes = "F2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(last, 2)}"
    dv_ok = DataValidation(type="list", formula1='"' + ",".join(TOPIC_OK) + '"', allow_blank=True)
    dv_dec = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    wb.create_sheet("Списки").sheet_state = "hidden"
    lst = wb["Списки"]
    for k, t in enumerate(topics, 1):
        lst.cell(row=k, column=1, value=t)
    dv_topic = DataValidation(type="list", formula1=f"=Списки!$A$1:$A${len(topics)}", allow_blank=True)
    for dv, col in ((dv_ok, "K"), (dv_topic, "L"), (dv_dec, "M")):
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(last, 2)}")
    ws.cell(row=last + 2, column=2, value=DISCLAIMER).font = f_bold
    wb.active = wb.sheetnames.index(SHEET_ROWS)
    wb.save(path)
    return len(rows)


def load_labels(path: str) -> list[dict]:
    """Прочитать заполненную анкету: по строке — ID правок, договор, тема и решение юриста."""
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_ROWS]
    roles = {}
    if SHEET_CONTRACTS in wb.sheetnames:
        for row in wb[SHEET_CONTRACTS].iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[4]:
                roles[str(row[0])] = str(row[4]).strip()
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1] or not row[14]:
            continue
        ok = (row[10] or "").strip()
        topic_prog = (row[9] or "").strip()
        topic_right = (row[11] or "").strip() or (topic_prog if ok == "Да" else "")
        out.append({
            "n": row[0], "contract": row[1], "role": roles.get(str(row[1]), ""), "clause": row[2] or "",
            "ids": [x.strip() for x in str(row[14]).split(",") if x.strip()],
            "topic_program": topic_prog, "topic_ok": ok, "topic": topic_right,
            "decision": (row[12] or "").strip(), "note": (row[13] or "").strip(),
        })
    return out
